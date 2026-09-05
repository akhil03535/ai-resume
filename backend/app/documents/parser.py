import io
import os

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from fastapi import UploadFile

from app.core.config import settings
from app.common.exceptions import FileProcessingError

# Magic byte signatures - never trust the filename extension alone (spec #9).
_PDF_MAGIC = b"%PDF-"
_DOCX_MAGIC = b"PK\x03\x04"  # docx is a zip archive


def validate_upload(file: UploadFile, raw_bytes: bytes) -> tuple[str, str]:
    """Validate extension, magic bytes, and size. Returns (extension, file_type)."""
    filename = file.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext not in settings.ALLOWED_UPLOAD_EXTENSIONS:
        raise FileProcessingError(f"Unsupported file type '{ext}'. Please upload a PDF or DOCX resume.")

    size_mb = len(raw_bytes) / (1024 * 1024)
    if size_mb > settings.MAX_UPLOAD_SIZE_MB:
        raise FileProcessingError(f"File is too large. Maximum size is {settings.MAX_UPLOAD_SIZE_MB}MB.")

    if ext == ".pdf" and not raw_bytes.startswith(_PDF_MAGIC):
        raise FileProcessingError("This file doesn't look like a valid PDF. It may be corrupted.")
    if ext == ".docx" and not raw_bytes.startswith(_DOCX_MAGIC):
        raise FileProcessingError("This file doesn't look like a valid DOCX. It may be corrupted.")

    file_type = "pdf" if ext == ".pdf" else "docx"
    return ext, file_type


def extract_text(file_path: str, file_type: str) -> str:
    try:
        with open(file_path, "rb") as file:
            return extract_text_from_bytes(file.read(), file_type)
    except FileProcessingError:
        raise
    except Exception as exc:
        raise FileProcessingError(
            "We couldn't read this file. It may be corrupted, scanned as an image, or password-protected."
        ) from exc


def extract_text_from_bytes(raw_bytes: bytes, file_type: str) -> str:
    try:
        if file_type == "pdf":
            return _extract_pdf_text(raw_bytes)
        elif file_type == "docx":
            return _extract_docx_text(raw_bytes)
        raise FileProcessingError(f"Unsupported file type: {file_type}")
    except FileProcessingError:
        raise
    except Exception as exc:
        raise FileProcessingError(
            "We couldn't read this file. It may be corrupted, scanned as an image, or password-protected."
        ) from exc


def _extract_pdf_text(raw_bytes: bytes) -> str:
    text_parts = []
    with fitz.open(stream=raw_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    text = "\n".join(text_parts).strip()
    if not text:
        raise FileProcessingError(
            "No extractable text was found in this PDF. It may be a scanned image - please upload a text-based PDF."
        )
    return text


def _extract_docx_text(raw_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(raw_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text = "\n".join(parts).strip()
    if not text:
        raise FileProcessingError("No extractable text was found in this DOCX file.")
    return text
