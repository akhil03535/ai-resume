import io

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x17, 0x25, 0x54)
SLATE = RGBColor(0x47, 0x55, 0x69)


def _heading(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def render_resume_docx(content: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    personal = content.get("personal_information", {})
    if personal.get("full_name"):
        p = doc.add_paragraph()
        run = p.add_run(personal["full_name"])
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = NAVY

    if personal.get("headline"):
        p = doc.add_paragraph(personal["headline"])
        p.runs[0].font.color.rgb = SLATE

    contact_bits = [v for v in [personal.get("email"), personal.get("phone"), personal.get("location")] if v]
    if contact_bits:
        p = doc.add_paragraph(" | ".join(contact_bits))
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = SLATE

    if content.get("summary"):
        _heading(doc, "Summary")
        doc.add_paragraph(content["summary"])

    if content.get("skills"):
        _heading(doc, "Skills")
        doc.add_paragraph(", ".join(content["skills"]))

    if content.get("experience"):
        _heading(doc, "Experience")
        for e in content["experience"]:
            p = doc.add_paragraph()
            p.add_run(f"{e.get('title', '')}, {e.get('company', '')}").bold = True
            dates = f"{e.get('start_date') or ''} - {'Present' if e.get('is_current') else (e.get('end_date') or '')}"
            doc.add_paragraph(dates).runs[0].font.color.rgb = SLATE
            for bullet in e.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if content.get("projects"):
        _heading(doc, "Projects")
        for proj in content["projects"]:
            p = doc.add_paragraph()
            p.add_run(proj.get("name", "")).bold = True
            for bullet in proj.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    if content.get("education"):
        _heading(doc, "Education")
        for ed in content["education"]:
            line = ed.get("degree") or ""
            if ed.get("field_of_study"):
                line += f", {ed['field_of_study']}"
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            doc.add_paragraph(ed.get("institution", "")).runs[0].font.color.rgb = SLATE

    if content.get("certifications"):
        _heading(doc, "Certifications")
        for c in content["certifications"]:
            label = c.get("name", "")
            if c.get("issuer"):
                label += f" — {c['issuer']}"
            doc.add_paragraph(label, style="List Bullet")

    if content.get("achievements"):
        _heading(doc, "Achievements")
        for a in content["achievements"]:
            label = a.get("title", "")
            if a.get("description"):
                label += f" — {a['description']}"
            doc.add_paragraph(label, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
